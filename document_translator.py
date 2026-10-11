from __future__ import annotations

import hashlib
import json
import logging
import os
import random
import re
import time
from pathlib import Path
from typing import Any

import fitz  # PyMuPDF
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from providers import build_provider
from translator import (
    QuotaExhaustedError,
    RetryExhaustedError,
    _retry_after_seconds,
    is_daily_quota_error,
    is_rate_limit_error,
    is_transient_api_error,
    _status_code,
)

log = logging.getLogger("document-translator")

SYSTEM_PROMPT = """You are an expert English-to-Persian technical document translator.
Translate English technical documentation into accurate, fluent Persian suitable for datasheets, reference manuals, and engineering books.
Return ONLY a valid JSON array. Each item must contain exactly the keys `id` and `translation`.
Rules:
- Translate only the supplied target items and return each id exactly once.
- Preserve meaning, technical detail, units, equations, numeric values, register addresses, pin names, identifiers, file paths, commands, code, URLs, product numbers, acronyms, and API names.
- Never alter protected placeholders such as ZXQPROTECT...QXZ; copy them exactly so the program can restore the original tokens.
- Do not translate identifiers such as STM32F103C8T6, USART1, GPIOA, HAL_GPIO_Init, 0x40021000, I2C, or section numbering.
- Use Persian digits only for ordinary Persian prose when appropriate. Never localize digits inside identifiers, hexadecimal values, part numbers, register names, model numbers, or numbered heading prefixes.
- Use glossary terms consistently. Do not add explanations or omit details.
- Keep the structure of each item; paragraph text should remain a paragraph, and table-cell text should remain concise where appropriate.
- The input may include fragments, headings, labels, captions, and table cells. Translate according to their local meaning.
"""


_TECH_TOKEN_RE = re.compile(
    r"(?<![A-Za-z0-9_])(?:"
    r"0[xX][0-9A-Fa-f]+|"
    r"[0-9]+(?:\.[0-9]+)?\s?(?:mV|V|uV|µV|mA|A|uA|µA|mW|W|kHz|MHz|GHz|Hz|ns|us|µs|ms|s|bit|bits|KB|MB|GB|kΩ|MΩ|Ω|%)|"
    r"[A-Za-z_][A-Za-z0-9_]*(?:[./:+-][A-Za-z0-9_]+)*|"
    r"[0-9]+(?:\.[0-9]+)*(?:[-/][A-Za-z0-9]+)+|"
    r"[A-Z]{2,}[A-Z0-9_]*"
    r")(?![A-Za-z0-9_])"
)
_HEADING_NUMBER_RE = re.compile(r"^(\s*(?:\d+\.)*\d+(?:\.)?)(\s+)")
_SENTENCE_END_RE = re.compile(r"[.!?؟؛:]\s*$")


def _looks_like_diagram_label(text: str, in_table: bool = False) -> bool:
    """Conservative heuristic for short vector-diagram labels extracted as text.

    Actual embedded images are never sent to the model. PDF converters sometimes turn
    vector labels into ordinary text; skip only short, label-like fragments outside tables.
    """
    value = " ".join(text.split())
    if in_table or not value or len(value) > 42 or _SENTENCE_END_RE.search(value):
        return False
    if len(value.split()) > 5:
        return False
    has_code = bool(re.search(r"[A-Z]{2,}\d|[A-Za-z]+\d|\d+[A-Za-z]|[_=<>→←↑↓]", value))
    has_diagram_symbol = bool(re.search(r"[_=<>→←↑↓|\[\]{}]", value))
    return has_code and (has_diagram_symbol or len(value.split()) <= 2)


def _protect_technical_tokens(text: str, item_id: str) -> tuple[str, dict[str, str]]:
    """Replace technical identifiers and leading heading numbering with stable placeholders."""
    replacements: dict[str, str] = {}
    counter = 0

    def stash(value: str) -> str:
        nonlocal counter
        token = f"ZXQPROTECT{item_id.upper()}X{counter}QXZ"
        counter += 1
        replacements[token] = value
        return token

    # Protect mixed letter/digit identifiers, hex values, API names, file paths, and code tokens.
    def replace_token(match):
        value = match.group(0)
        # Pure alphabetic words are translated normally unless they are all-caps acronyms.
        if not re.search(r"\d|[_./:+-]", value) and not value.isupper():
            return value
        return stash(value)

    text = _TECH_TOKEN_RE.sub(replace_token, text)

    # Keep numbering such as 1.2.3 in English digits and preserve its exact punctuation.
    # Do this after token protection so the placeholder itself is not protected a second time.
    heading_match = _HEADING_NUMBER_RE.match(text)
    if heading_match:
        prefix = heading_match.group(1)
        whitespace = heading_match.group(2)
        protected_prefix = stash(prefix)
        text = protected_prefix + whitespace + text[heading_match.end():]

    return text, replacements


def _restore_protected_tokens(text: str, replacements: dict[str, str]) -> str:
    for token, original in replacements.items():
        text = text.replace(token, original)
    return text




def flatten_glossary(glossary: dict) -> dict[str, str]:
    result: dict[str, str] = {}
    categories = glossary.get("categories", glossary) if isinstance(glossary, dict) else {}
    for category_terms in categories.values():
        if isinstance(category_terms, dict):
            for term, translation in category_terms.items():
                if isinstance(term, str) and isinstance(translation, str):
                    result[term] = translation
    return result


def relevant_glossary(glossary: dict, texts: list[str], limit: int = 100) -> dict[str, str]:
    terms = flatten_glossary(glossary)
    corpus = "\n".join(texts).lower()
    matched = [(len(k), k, v) for k, v in terms.items() if k.lower() in corpus]
    matched.sort(key=lambda x: (-x[0], x[1].lower()))
    return {k: v for _, k, v in matched[:limit]}


def parse_json_array(content: str) -> list[dict]:
    text = content.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.I)
    start, end = text.find("["), text.rfind("]")
    if start < 0 or end < start:
        raise ValueError("Model response did not contain a JSON array")
    data = json.loads(text[start:end + 1])
    if not isinstance(data, list):
        raise ValueError("Model response JSON is not an array")
    return data


def validate_result(raw: str, batch: list[dict]) -> dict[str, str]:
    items = parse_json_array(raw)
    expected = {x["id"] for x in batch}
    result = {}
    for item in items:
        if not isinstance(item, dict) or set(("id", "translation")) - set(item):
            raise ValueError("Invalid translated item schema")
        item_id = str(item["id"])
        translation = item["translation"]
        if item_id not in expected or item_id in result or not isinstance(translation, str) or not translation.strip():
            raise ValueError("Unexpected, duplicate, or empty translation item")
        result[item_id] = translation.strip()
    if set(result) != expected:
        raise ValueError(f"Translation item mismatch: expected {len(expected)}, got {len(result)}")
    return result


def api_call(provider, batch, glossary, cfg, all_targets):
    t = cfg.get("translation", {})
    index_by_id = {item["id"]: i for i, item in enumerate(all_targets)}
    first_i = index_by_id[batch[0]["id"]]
    last_i = index_by_id[batch[-1]["id"]]
    context_n = max(0, int(t.get("document_context_items", 2)))
    previous_context = all_targets[max(0, first_i - context_n):first_i]
    following_context = all_targets[last_i + 1:last_i + 1 + context_n]
    corpus = [x["text"] for x in batch + previous_context + following_context]
    protected = [term for term in glossary.get("do_not_translate", [])
                 if isinstance(term, str) and term.lower() in "\n".join(corpus).lower()]
    protect_maps = {}
    protected_batch = []
    for item in batch:
        protected_text, replacements = _protect_technical_tokens(item["text"], item["id"])
        protect_maps[item["id"]] = replacements
        protected_item = dict(item)
        protected_item["text"] = protected_text
        protected_batch.append(protected_item)
    payload = {
        "task": "Translate the target technical document paragraphs and table cells from English to Persian.",
        "style": t.get("document_style", "Technically precise Persian for embedded systems, electronics, firmware, and engineering books."),
        "glossary": relevant_glossary(glossary, corpus, int(t.get("document_glossary_max_terms", 100))),
        "protected_terms": protected,
        "previous_context": [{"kind": x["kind"], "text": x["text"]} for x in previous_context],
        "following_context": [{"kind": x["kind"], "text": x["text"]} for x in following_context],
        "target": [{"id": x["id"], "kind": x["kind"], "text": x["text"]} for x in protected_batch],
    }
    retries = max(0, int(t.get("retry_count", 5)))
    base = max(0.1, float(t.get("retry_base_seconds", 10)))
    max_delay = max(base, float(t.get("retry_max_seconds", 180)))
    jitter = max(0.0, float(t.get("retry_jitter_seconds", 3)))
    last = None
    for attempt in range(retries + 1):
        try:
            response = provider.generate(SYSTEM_PROMPT, json.dumps(payload, ensure_ascii=False), temperature=float(t.get("temperature", 0.15)), max_output_tokens=int(t.get("document_max_output_tokens", 6000)))
            translated = validate_result(response.content, batch)
            return {item_id: _restore_protected_tokens(value, protect_maps.get(item_id, {})) for item_id, value in translated.items()}
        except Exception as exc:
            last = exc
            if is_daily_quota_error(exc):
                raise QuotaExhaustedError("Provider daily/project quota is exhausted.", last_error=exc, retry_after_seconds=_retry_after_seconds(exc)) from exc
            # Invalid JSON/schema is not a transient HTTP failure; caller can split the batch.
            if isinstance(exc, (ValueError, json.JSONDecodeError)):
                raise
            if not is_transient_api_error(exc):
                raise
            if attempt >= retries:
                break
            provider_delay = _retry_after_seconds(exc) if is_rate_limit_error(exc) else None
            delay = min(max_delay, max(base, provider_delay)) if provider_delay is not None else min(max_delay, base * (2 ** attempt)) + random.uniform(0, jitter)
            log.warning("Temporary API error (%s): %s; retry %d/%d in %.1fs", _status_code(exc), exc, attempt + 1, retries, delay)
            time.sleep(delay)
    raise RetryExhaustedError(f"Document translation request failed after retries: {last}", last_error=last) from last




_PAGE_NUMBER_RE = re.compile(r"^(?:(?:page|p\.?|صفحه)\s*)?[-–—]?\s*\d{1,4}\s*(?:[-–—]|/\s*\d{1,4})?$", re.I)


def _normalise_block_text(raw: str) -> str:
    """Turn PDF line wraps inside one text block into a readable paragraph."""
    lines = [re.sub(r"\s+", " ", line).strip() for line in raw.replace("\r", "\n").split("\n")]
    lines = [line for line in lines if line]
    if not lines:
        return ""
    result = lines[0]
    for line in lines[1:]:
        # Rejoin words broken by a PDF line-wrap hyphen; otherwise add a space.
        if result.endswith("-") and line and line[0].islower():
            result = result[:-1] + line
        else:
            result += " " + line
    return re.sub(r"\s+", " ", result).strip()


def _norm_repeat_key(text: str) -> str:
    text = text.lower()
    text = re.sub(r"\d+", "#", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip(" -–—|·•")


def extract_pdf_targets(pdf_path: Path) -> list[dict[str, Any]]:
    """Extract text blocks only. Images, logos, page numbers and repeated headers/footers are omitted."""
    pdf = fitz.open(str(pdf_path))
    page_records: list[list[dict[str, Any]]] = []
    top_bottom_occurrences: dict[str, set[int]] = {}
    page_count = len(pdf)

    for page_index, page in enumerate(pdf):
        page_rect = page.rect
        records = []
        for block in page.get_text("blocks", sort=True):
            if len(block) < 7 or int(block[6]) != 0:
                continue
            x0, y0, x1, y1, raw = block[:5]
            text = _normalise_block_text(raw)
            if not text:
                continue
            # Isolated page-number blocks are never useful content.
            if _PAGE_NUMBER_RE.fullmatch(text):
                continue
            # Omit text placed in the extreme top/bottom page edge, where logos and page furniture live.
            if y0 < page_rect.height * 0.045 or y1 > page_rect.height * 0.955:
                continue
            # Very small isolated glyph fragments are usually layout artifacts.
            if len(text) <= 2 and not re.search(r"[A-Za-z0-9]", text):
                continue
            key = _norm_repeat_key(text)
            in_margin = y0 < page_rect.height * 0.13 or y1 > page_rect.height * 0.87
            if in_margin and key and len(key) <= 180:
                top_bottom_occurrences.setdefault(key, set()).add(page_index)
            records.append({"page": page_index, "bbox": (x0, y0, x1, y1), "text": text, "key": key})
        page_records.append(records)
    pdf.close()

    # Repeated text in the top/bottom margins is treated as running header/footer.
    repeat_threshold = max(2, int(page_count * 0.20))
    repeated_margin_keys = {key for key, pages in top_bottom_occurrences.items() if len(pages) >= repeat_threshold}

    targets: list[dict[str, Any]] = []
    for records in page_records:
        for record in records:
            if record["key"] in repeated_margin_keys:
                continue
            text = record["text"]
            # Avoid obvious image/diagram labels that are just a short technical token.
            # This does not affect the translation engine; it only excludes likely diagram fragments.
            if len(text) <= 24 and len(text.split()) <= 3 and re.fullmatch(r"[A-Z0-9_./:+-]+", text):
                continue
            targets.append({"id": f"p{len(targets)+1}", "kind": "paragraph", "location": f"page-{record['page']+1}/block-{len(targets)+1}", "text": text})
    return targets


def extract_docx_targets(docx_path: Path) -> list[dict[str, Any]]:
    """Optional DOCX input: read body paragraphs only; ignore tables, headers, footers and images."""
    doc = Document(str(docx_path))
    targets = []
    for i, para in enumerate(doc.paragraphs, start=1):
        text = re.sub(r"\s+", " ", para.text).strip()
        if not text or _PAGE_NUMBER_RE.fullmatch(text):
            continue
        if para._p.xpath('.//w:drawing') or para._p.xpath('.//w:pict'):
            continue
        targets.append({"id": f"p{len(targets)+1}", "kind": "paragraph", "location": f"paragraph-{i}", "text": text})
    return targets


def extract_targets(input_path: Path) -> list[dict[str, Any]]:
    suffix = input_path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf_targets(input_path)
    if suffix == ".docx":
        return extract_docx_targets(input_path)
    raise ValueError("Supported input formats are PDF and DOCX.")


def _make_clean_docx(translated_items: list[dict[str, Any]], output_path: Path) -> None:
    """Build a new, plain Word file from scratch: no source formatting, images, headers or footers."""
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for item in translated_items:
        text = item["translation"].strip()
        if not text:
            continue
        paragraph = doc.add_paragraph(style="Normal")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        paragraph.paragraph_format.keep_together = False
        paragraph.paragraph_format.keep_with_next = False
        run = paragraph.add_run(text)
        # Keep technical identifiers and digits intact; do not force language-specific digit conversion.
        run.font.name = "Calibri"
        run.font.size = Pt(11)
        p_pr = paragraph._p.get_or_add_pPr()
        bidi = OxmlElement("w:bidi")
        bidi.set(qn("w:val"), "1")
        p_pr.append(bidi)
        r_pr = run._r.get_or_add_rPr()
        rtl = OxmlElement("w:rtl")
        rtl.set(qn("w:val"), "1")
        r_pr.append(rtl)

    # New document has no source header/footer, logos, page numbering, tables, or embedded images.
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))


def translate_document(input_path: Path, output_path: Path, cfg: dict, provider_name: str, resume: bool = False):
    checkpoint_path = Path(str(output_path) + ".checkpoint.json")
    targets = extract_targets(input_path)
    log.info("Extracted %d text paragraphs. Output will be rebuilt as plain Word text.", len(targets))
    if not targets:
        raise RuntimeError("No selectable text found. If this is a scanned/image-only PDF, OCR is required and is not included in this version.")

    digest = hashlib.sha256(input_path.read_bytes()).hexdigest()
    old: dict[str, str] = {}
    if checkpoint_path.exists():
        state = json.loads(checkpoint_path.read_text(encoding="utf-8"))
        if state.get("input_sha256") != digest:
            raise RuntimeError("Input differs from checkpoint. Choose a new output path or remove the old checkpoint.")
        if resume:
            old = state.get("translations", {})

    provider = build_provider(cfg, provider_name)
    glossary_path = Path(cfg.get("glossary_file", "glossary.json"))
    if not glossary_path.is_absolute():
        glossary_path = Path.cwd() / glossary_path
    glossary = json.loads(glossary_path.read_text(encoding="utf-8")) if glossary_path.exists() else {}

    settings = cfg.get("translation", {})
    max_items = max(1, int(settings.get("document_batch_max_items", 10)))
    max_chars = max(500, int(settings.get("document_batch_max_chars", 6000)))
    translations = dict(old)
    pending = [item for item in targets if item["id"] not in translations]
    log.info("Already translated: %d/%d", len(translations), len(targets))

    def commit_state(status: str, current: list[str] | None = None, error: str | None = None):
        state = {
            "version": 2, "status": status, "input": str(input_path.resolve()),
            "input_sha256": digest, "output": str(output_path.resolve()),
            "provider": provider_name,
            "model": cfg.get("providers", {}).get(provider_name, {}).get("model"),
            "completed_count": len(translations), "total_count": len(targets),
            "translations": translations, "current_batch": current, "last_error": error,
        }
        tmp = checkpoint_path.with_suffix(checkpoint_path.suffix + ".tmp")
        tmp.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")
        os.replace(tmp, checkpoint_path)

    commit_state("translating")

    def translate_batch_resilient(batch: list[dict[str, Any]]):
        ids = [item["id"] for item in batch]
        try:
            result = api_call(provider, batch, glossary, cfg, targets)
            translations.update(result)
            commit_state("translating")
            log.info("Translated %d/%d paragraphs", len(translations), len(targets))
        except (ValueError, json.JSONDecodeError) as exc:
            if len(batch) <= 1:
                raise
            log.warning("Invalid model output for batch %s; splitting batch. Reason: %s", ids, exc)
            mid = len(batch) // 2
            translate_batch_resilient(batch[:mid])
            translate_batch_resilient(batch[mid:])
        except (QuotaExhaustedError, RetryExhaustedError) as exc:
            commit_state("paused", ids, str(exc))
            raise

    index = 0
    while index < len(pending):
        batch: list[dict[str, Any]] = []
        char_count = 0
        while index < len(pending) and len(batch) < max_items:
            item = pending[index]
            if batch and char_count + len(item["text"]) > max_chars:
                break
            batch.append(item)
            char_count += len(item["text"])
            index += 1
        translate_batch_resilient(batch)

    ordered = [{"id": item["id"], "translation": translations[item["id"]]} for item in targets]
    _make_clean_docx(ordered, output_path)
    commit_state("completed")
    log.info("Saved plain Word document: %s", output_path)

